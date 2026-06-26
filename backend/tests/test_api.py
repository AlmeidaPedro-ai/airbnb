"""Testes de integração da API (Fase 2).

Requerem um Postgres migrado e populado (seed) apontado por DATABASE_URL.
Se o banco não estiver acessível, o módulo é ignorado (skip).
"""
from __future__ import annotations

import json
import os

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import text

pytestmark = pytest.mark.api


@pytest.fixture(scope="module")
def client():
    from app.database import engine
    from app.main import app

    try:
        with engine.connect() as conn:
            conn.execute(text("SELECT 1"))
    except Exception as exc:  # pragma: no cover - ambiente sem banco
        pytest.skip(f"Banco indisponível para testes de API: {exc}")
    return TestClient(app)


@pytest.fixture(scope="module")
def token(client):
    senha = os.getenv("SEED_USER_PASSWORD", "admin123")
    email = os.getenv("SEED_USER_EMAIL", "1994.pedro@gmail.com")
    resp = client.post("/api/auth/login", json={"email": email, "senha": senha})
    assert resp.status_code == 200, resp.text
    return resp.json()["access_token"]


@pytest.fixture
def auth(token):
    return {"Authorization": f"Bearer {token}"}


# --------------------------------------------------------------------------- #
# Autenticação                                                                 #
# --------------------------------------------------------------------------- #
def test_login_invalido(client):
    resp = client.post(
        "/api/auth/login", json={"email": "x@y.com", "senha": "errada"}
    )
    assert resp.status_code == 401


def test_rota_protegida_sem_token(client):
    assert client.get("/api/apartamentos").status_code == 401


def test_me(client, auth):
    resp = client.get("/api/auth/me", headers=auth)
    assert resp.status_code == 200
    assert "email" in resp.json()


# --------------------------------------------------------------------------- #
# Métricas — paridade com a planilha (leitura, roda antes das mutações)        #
# --------------------------------------------------------------------------- #
def test_metricas_todos_aceite(client, auth):
    resp = client.get(
        "/api/metricas",
        params={"inicio": "2026-01-01", "fim": "2026-03-31"},
        headers=auth,
    )
    assert resp.status_code == 200, resp.text
    m = resp.json()
    assert m["num_reservas"] == 15
    assert m["noites_reservadas"] == 69
    assert m["total_hospedes"] == 39
    assert m["receita_diarias"] == 22300.0
    assert m["receita_liquida_recebida"] == 23958.0
    assert m["adr"] == 323.19
    assert m["noites_disponiveis"] == 270
    assert m["ocupacao"] == 0.2556
    assert m["revpar"] == 82.59
    assert m["despesas_totais"] == 8510.0
    assert m["imposto_estimado"] == 1541.0
    assert m["lucro_liquido"] == 13907.0


def test_metricas_por_apartamento(client, auth):
    resp = client.get(
        "/api/metricas/por-apartamento",
        params={"inicio": "2026-01-01", "fim": "2026-03-31"},
        headers=auth,
    )
    assert resp.status_code == 200
    por = {x["nome"]: x for x in resp.json()}
    assert por["Copacabana 302"]["noites"] == 25
    assert por["Copacabana 302"]["adr"] == 300.0
    assert por["Ipanema 1104"]["noites"] == 30
    assert por["Ipanema 1104"]["despesas"] == 4790.0
    assert por["Botafogo 506"]["noites"] == 14


def test_imposto_aceite(client, auth):
    resp = client.get("/api/imposto", params={"ano": 2026}, headers=auth)
    assert resp.status_code == 200, resp.text
    g = resp.json()
    assert g["total_devido"] == 1541.0
    jan, fev, mar = g["linhas"][0], g["linhas"][1], g["linhas"][2]
    assert jan["imposto_devido"] == 685.25 and jan["isento"] is False
    assert fev["imposto_devido"] == 855.75
    assert mar["isento"] is True and mar["imposto_devido"] == 0.0
    assert "contador" in g["disclaimer"]


def test_metricas_mensal(client, auth):
    resp = client.get(
        "/api/metricas/mensal",
        params={"inicio": "2026-01-01", "fim": "2026-03-31"},
        headers=auth,
    )
    assert resp.status_code == 200, resp.text
    serie = resp.json()
    assert [m["rotulo"] for m in serie] == ["01/2026", "02/2026", "03/2026"]
    assert serie[0]["receita_diarias"] == 7600.0
    assert serie[0]["imposto_estimado"] == 685.25
    assert serie[2]["imposto_estimado"] == 0.0


def test_imposto_calcular_outras_rendas(client, auth):
    resp = client.post(
        "/api/imposto/calcular",
        json={"ano": 2026, "outras_rendas": {"3": 2000}},
        headers=auth,
    )
    assert resp.status_code == 200, resp.text
    g = resp.json()
    mar = g["linhas"][2]
    # Mar: base 4305 + 2000 = 6305 -> 6305*0.275 - 896 = 837.88
    assert mar["base"] == 6305.0
    assert mar["imposto_devido"] == 837.88
    assert g["total_devido"] == 2378.88


def test_metricas_periodo_invalido(client, auth):
    resp = client.get(
        "/api/metricas",
        params={"inicio": "2026-03-31", "fim": "2026-01-01"},
        headers=auth,
    )
    assert resp.status_code == 422


def test_parametros_ir(client, auth):
    resp = client.get("/api/parametros-ir", headers=auth)
    assert resp.status_code == 200
    p = resp.json()
    assert p["ano_vigencia"] == 2026
    assert len(p["faixas"]) == 5
    assert p["aplicar_redutor_transicao"] is False


# --------------------------------------------------------------------------- #
# CRUD e regras de negócio (mutações isoladas, com limpeza)                    #
# --------------------------------------------------------------------------- #
def test_apartamento_crud_e_delete_conflito(client, auth):
    # Cria apartamento INATIVO (não afeta contagem de apês ativos nas métricas).
    novo = {
        "nome": "Teste CRUD 999",
        "endereco": "Rua Teste, 1",
        "quartos": 1,
        "capacidade": 2,
        "data_inicio_operacao": "2026-01-01",
        "ativo": False,
    }
    r = client.post("/api/apartamentos", json=novo, headers=auth)
    assert r.status_code == 201, r.text
    apto_id = r.json()["id"]

    # Nome duplicado -> 409
    assert client.post("/api/apartamentos", json=novo, headers=auth).status_code == 409

    # Update
    r = client.put(
        f"/api/apartamentos/{apto_id}", json={"capacidade": 5}, headers=auth
    )
    assert r.status_code == 200 and r.json()["capacidade"] == 5

    # Apartamento seedado (id=1) tem reservas -> delete 409
    r = client.delete("/api/apartamentos/1", headers=auth)
    assert r.status_code == 409

    # Delete do apartamento de teste (sem vínculos) -> 204
    assert client.delete(f"/api/apartamentos/{apto_id}", headers=auth).status_code == 204
    assert client.get(f"/api/apartamentos/{apto_id}", headers=auth).status_code == 404


def test_reserva_avisos_capacidade_e_sobreposicao(client, auth):
    # Apartamento inativo dedicado ao teste.
    r = client.post(
        "/api/apartamentos",
        json={
            "nome": "Teste Avisos 998",
            "quartos": 1,
            "capacidade": 2,
            "data_inicio_operacao": "2026-01-01",
            "ativo": False,
        },
        headers=auth,
    )
    apto_id = r.json()["id"]
    try:
        base = {
            "apartamento_id": apto_id,
            "check_in": "2026-05-01",
            "check_out": "2026-05-05",
            "num_hospedes": 2,
            "valor_hospedagem": "1000.00",
            "taxa_limpeza": "100.00",
            "comissao": "40.00",
        }
        r1 = client.post("/api/reservas", json=base, headers=auth)
        assert r1.status_code == 201, r1.text
        res1_id = r1.json()["reserva"]["id"]
        assert r1.json()["reserva"]["noites"] == 4
        assert r1.json()["reserva"]["receita_liquida"] == "1060.00"

        # Sobreposta + capacidade excedida -> 2 avisos.
        sobreposta = {
            **base,
            "check_in": "2026-05-03",
            "check_out": "2026-05-08",
            "num_hospedes": 4,
        }
        r2 = client.post("/api/reservas", json=sobreposta, headers=auth)
        assert r2.status_code == 201
        avisos = r2.json()["avisos"]
        assert any("capacidade" in a.lower() for a in avisos)
        assert any("sobreposi" in a.lower() for a in avisos)
        res2_id = r2.json()["reserva"]["id"]

        # check_out <= check_in -> 422 (validação Pydantic)
        ruim = {**base, "check_in": "2026-05-10", "check_out": "2026-05-10"}
        assert client.post("/api/reservas", json=ruim, headers=auth).status_code == 422

        # Limpeza
        assert client.delete(f"/api/reservas/{res2_id}", headers=auth).status_code == 204
        assert client.delete(f"/api/reservas/{res1_id}", headers=auth).status_code == 204
    finally:
        client.delete(f"/api/apartamentos/{apto_id}", headers=auth)


def test_importacao_reservas_fluxo(client, auth):
    # Apartamento inativo dedicado para não afetar métricas; com reservas
    # importadas em datas futuras (fora dos testes de aceite).
    r = client.post(
        "/api/apartamentos",
        json={
            "nome": "Import Teste 997",
            "quartos": 1,
            "capacidade": 4,
            "data_inicio_operacao": "2026-01-01",
            "ativo": False,
        },
        headers=auth,
    )
    apto_id = r.json()["id"]
    csv_texto = (
        "Confirmation code,Listing,Start date,End date,# of guests,Earnings\n"
        "HMQA01,Import Teste,08/08/2026,12/08/2026,2,\"1.200,00\"\n"
        "HMQA02,Import Teste,15/08/2026,18/08/2026,3,\"900,00\"\n"
        "HMQA01,Import Teste,20/08/2026,22/08/2026,2,\"600,00\"\n"  # duplicada
    )
    arquivo = ("reservas.csv", csv_texto, "text/csv")
    try:
        # 1. Analisar
        r = client.post(
            "/api/importar/reservas/analisar",
            files={"arquivo": arquivo},
            headers=auth,
        )
        assert r.status_code == 200, r.text
        analise = r.json()
        assert analise["sugestao_mapeamento"]["check_in"] == "Start date"
        assert analise["listings"] == ["Import Teste"]
        assert analise["total_linhas"] == 3

        mapeamento = json.dumps(analise["sugestao_mapeamento"])
        de_para = json.dumps({"Import Teste": apto_id})

        # 2. Dry-run
        r = client.post(
            "/api/importar/reservas/confirmar",
            files={"arquivo": arquivo},
            data={"mapeamento": mapeamento, "de_para": de_para, "dry_run": "true"},
            headers=auth,
        )
        assert r.status_code == 200, r.text
        previa = r.json()
        assert previa["criar"] == 2
        assert previa["ignorar"] == 1  # código duplicado
        assert previa["importadas"] == 0

        # 3. Importar de fato
        r = client.post(
            "/api/importar/reservas/confirmar",
            files={"arquivo": arquivo},
            data={"mapeamento": mapeamento, "de_para": de_para, "dry_run": "false"},
            headers=auth,
        )
        assert r.json()["importadas"] == 2

        # 4. Reimportar -> tudo deduplicado (0 criadas)
        r = client.post(
            "/api/importar/reservas/confirmar",
            files={"arquivo": arquivo},
            data={"mapeamento": mapeamento, "de_para": de_para, "dry_run": "false"},
            headers=auth,
        )
        assert r.json()["importadas"] == 0
        assert r.json()["criar"] == 0
    finally:
        # Limpa as reservas importadas e o apartamento.
        reservas = client.get(
            "/api/reservas", params={"apartamento_id": apto_id}, headers=auth
        ).json()
        for rv in reservas:
            client.delete(f"/api/reservas/{rv['id']}", headers=auth)
        client.delete(f"/api/apartamentos/{apto_id}", headers=auth)


def test_exportacao_reservas_csv(client, auth):
    r = client.get(
        "/api/exportar/reservas",
        params={"formato": "csv", "inicio": "2026-01-01", "fim": "2026-03-31"},
        headers=auth,
    )
    assert r.status_code == 200
    assert "text/csv" in r.headers["content-type"]
    assert "attachment" in r.headers["content-disposition"]
    texto = r.content.decode("utf-8-sig")
    assert "Apartamento" in texto.splitlines()[0]
    assert "Copacabana 302" in texto


def test_exportacao_imposto_xlsx(client, auth):
    from openpyxl import load_workbook
    import io as _io

    r = client.get(
        "/api/exportar/imposto",
        params={"formato": "xlsx", "ano": 2026},
        headers=auth,
    )
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]
    wb = load_workbook(_io.BytesIO(r.content))
    ws = wb.active
    assert ws.cell(row=1, column=1).value == "Mês"
    # Última linha = total devido 1541.00
    ultima = ws.max_row
    assert ws.cell(row=ultima, column=1).value == "Total"
    assert float(ws.cell(row=ultima, column=10).value) == 1541.0


def test_exportacao_formato_invalido(client, auth):
    r = client.get(
        "/api/exportar/reservas", params={"formato": "pdf"}, headers=auth
    )
    assert r.status_code == 400


def test_despesa_crud(client, auth):
    r = client.post(
        "/api/despesas",
        json={
            "apartamento_id": None,
            "data": "2026-07-15",
            "categoria": "Internet",
            "valor": "130.00",
            "dedutivel_ir": False,
        },
        headers=auth,
    )
    assert r.status_code == 201, r.text
    desp_id = r.json()["id"]
    assert client.delete(f"/api/despesas/{desp_id}", headers=auth).status_code == 204
