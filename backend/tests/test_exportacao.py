"""Testes dos geradores de exportação CSV/XLSX."""
from __future__ import annotations

import io
from datetime import date
from decimal import Decimal

from app.services.exportacao import gerar_csv, gerar_xlsx

D = Decimal


def test_gerar_csv_basico():
    dados = gerar_csv(
        ["Data", "Valor", "Dedutível"],
        [[date(2026, 1, 31), D("650.00"), True]],
    )
    assert dados.startswith(b"\xef\xbb\xbf")  # BOM UTF-8
    texto = dados.decode("utf-8-sig")
    linhas = texto.strip().split("\r\n")
    assert linhas[0] == "Data;Valor;Dedutível"
    assert linhas[1] == "31/01/2026;650,00;Sim"  # vírgula decimal, data BR


def test_gerar_xlsx_reabre():
    from openpyxl import load_workbook

    dados = gerar_xlsx(
        ["Mês", "Imposto"],
        [["Jan", D("685.25")], ["Fev", D("855.75")]],
        nome_aba="Imposto 2026",
    )
    wb = load_workbook(io.BytesIO(dados))
    ws = wb.active
    assert ws.title == "Imposto 2026"
    assert [c.value for c in ws[1]] == ["Mês", "Imposto"]
    assert ws.cell(row=2, column=1).value == "Jan"
    assert float(ws.cell(row=2, column=2).value) == 685.25
    assert float(ws.cell(row=3, column=2).value) == 855.75
